"""hwp/hwpx/docx/xlsx 정규화 테스트용 합성 문서 생성 헬퍼."""

import io


def docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("1. Project Name : Docx Smoke Test Works")
    doc.add_paragraph("2. Contract Amount : KRW 220,000,000")
    doc.add_paragraph("3. Contract Date : 2026-01-15")
    doc.add_paragraph("7. Client (Creditor) : Docx Test Corporation")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "계약정보"
    ws.append(["2. Contract Amount : KRW 99,000,000"])
    ws.append(["3. Contract Date : 2026-02-20"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
