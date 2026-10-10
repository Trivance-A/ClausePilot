import io
import shutil
import struct
import zipfile
import zlib

import pymupdf as fitz
import pytest

from ai import parser
from ai.parser import (
    UnsupportedFormatError,
    blocks_to_lines,
    normalize_to_pdf,
    parse_docx,
    parse_hwpx,
    parse_xlsx,
)
from ai.parser.hwp import (
    HWPTAG_CTRL_HEADER,
    HWPTAG_LIST_HEADER,
    HWPTAG_PARA_HEADER,
    HWPTAG_PARA_TEXT,
    HWPTAG_TABLE,
    para_text,
    parse_section,
)
from ai.tests.fixtures import contract_pdf


def record(tag: int, level: int, payload: bytes) -> bytes:
    return struct.pack("<I", tag | (level << 10) | (len(payload) << 20)) + payload


def text(s: str) -> bytes:
    return s.encode("utf-16-le")


def test_hwp_para_text_skips_inline_controls():
    tab = struct.pack("<H", 9) + b"\x00" * 14  # inline control = 8 WCHARs
    table_ctrl = struct.pack("<H", 11) + b"\x00" * 14
    assert para_text(text("계약") + tab + text("금액") + table_ctrl + struct.pack("<H", 13)) == "계약\t금액\n"


def test_hwp_section_paragraphs_and_table_cells():
    def cell(row, col):
        return struct.pack("<HHI", 1, 0, 0) + struct.pack("<HHHH", col, row, 1, 1)

    data = b"".join(
        [
            record(HWPTAG_PARA_HEADER, 0, b""),
            record(HWPTAG_PARA_TEXT, 1, text("공사도급계약서")),
            record(HWPTAG_PARA_HEADER, 0, b""),
            record(HWPTAG_CTRL_HEADER, 1, b" lbt"),
            record(HWPTAG_TABLE, 2, b"\x00" * 8),
            record(HWPTAG_LIST_HEADER, 2, cell(0, 0)),
            record(HWPTAG_PARA_HEADER, 2, b""),
            record(HWPTAG_PARA_TEXT, 3, text("계약금액")),
            record(HWPTAG_LIST_HEADER, 2, cell(0, 1)),
            record(HWPTAG_PARA_HEADER, 2, b""),
            record(HWPTAG_PARA_TEXT, 3, text("금 삼억원정")),
            record(HWPTAG_PARA_HEADER, 0, b""),
            record(HWPTAG_PARA_TEXT, 1, text("제1조(목적)")),
        ]
    )
    blocks = parse_section(data)
    assert blocks == [
        {"type": "p", "text": "공사도급계약서"},
        {"type": "table", "cells": [[1, 1, "계약금액"], [1, 2, "금 삼억원정"]]},
        {"type": "p", "text": "제1조(목적)"},
    ]
    assert blocks_to_lines(blocks) == ["공사도급계약서", "계약금액 | 금 삼억원정", "제1조(목적)"]
    assert parse_section(zlib.decompress(zlib.compress(data)))  # raw records survive (de)compression


def hwpx_bytes() -> bytes:
    ns = 'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph" xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section"'
    cell = '<hp:tc><hp:subList><hp:p><hp:run><hp:t>{t}</hp:t></hp:run></hp:p></hp:subList><hp:cellAddr colAddr="{c}" rowAddr="{r}"/></hp:tc>'
    tbl = "<hp:tbl><hp:tr>" + cell.format(t="계약금액", c=0, r=0) + cell.format(t="금 삼억원정", c=1, r=0) + "</hp:tr></hp:tbl>"
    xml = f"<hs:sec {ns}><hp:p><hp:run><hp:t>공사명 : 한빛시 도로공사</hp:t></hp:run></hp:p><hp:p><hp:run>{tbl}</hp:run></hp:p></hs:sec>"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("Contents/section0.xml", xml)
    return buf.getvalue()


def test_hwpx(tmp_path):
    path = tmp_path / "a.hwpx"
    path.write_bytes(hwpx_bytes())
    assert parse_hwpx(str(path)) == [
        {"type": "p", "text": "공사명 : 한빛시 도로공사"},
        {"type": "table", "cells": [[1, 1, "계약금액"], [1, 2, "금 삼억원정"]]},
    ]


def test_docx_keeps_body_order_and_merged_cells(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("공사도급계약서")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "계약금액", "금 삼억원정"
    table.cell(1, 0).merge(table.cell(1, 1)).text = "특약사항 없음"
    doc.add_paragraph("제1조(목적)")
    path = tmp_path / "a.docx"
    doc.save(path)
    blocks = parse_docx(str(path))
    assert [b["type"] for b in blocks] == ["p", "table", "p"]
    assert blocks[1]["cells"] == [[1, 1, "계약금액"], [1, 2, "금 삼억원정"], [2, 1, "특약사항 없음"]]


def test_xlsx(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "견적"
    wb.active.append(["품목", "금액"])
    wb.active.append(["포장공사", 35000000])
    path = tmp_path / "a.xlsx"
    wb.save(path)
    assert blocks_to_lines(parse_xlsx(str(path))) == ["[견적]", "품목 | 금액", "포장공사 | 35000000"]


def test_unsupported_format(tmp_path):
    with pytest.raises(UnsupportedFormatError):
        normalize_to_pdf(str(tmp_path / "a.txt"), "txt", str(tmp_path / "o.pdf"))


def test_pdf_passthrough_meta_and_router(tmp_path):
    out = normalize_to_pdf(contract_pdf(str(tmp_path / "c.pdf")), "pdf", str(tmp_path / "n.pdf"))
    assert out["source_type"] == "native" and out["parsed_text_path"] is None
    assert out["pages"] == [{"page_no": 1, "width_pt": 595.0, "height_pt": 842.0, "rotation": 0, "has_text_layer": True}]


def test_hwpx_without_libreoffice_falls_back_to_text_pdf(tmp_path, monkeypatch):
    monkeypatch.setattr(parser, "libreoffice_to_pdf", lambda src, out: False)
    src = tmp_path / "a.hwpx"
    src.write_bytes(hwpx_bytes())
    out = normalize_to_pdf(str(src), "hwpx", str(tmp_path / "n.pdf"))
    assert out["parsed_text_path"] and (tmp_path / "n.parsed_text.json").exists()
    with fitz.open(out["normalized_pdf_path"]) as doc:
        assert "계약금액 | 금 삼억원정" in doc[0].get_text()


def test_broken_native_parse_without_libreoffice_fails_clearly(tmp_path, monkeypatch):
    monkeypatch.setattr(parser, "libreoffice_to_pdf", lambda src, out: False)
    src = tmp_path / "a.hwp"
    src.write_bytes(b"HWP Document File V3.00 \x1a\x01\x02\x03\x04\x05" + b"\x00" * 64)
    with pytest.raises(ValueError, match="HWP 3.0"):
        normalize_to_pdf(str(src), "hwp", str(tmp_path / "n.pdf"))


def test_png_becomes_scan_pdf(tmp_path):
    doc = fitz.open()
    doc.new_page(width=200, height=200)
    png = tmp_path / "a.png"
    doc[0].get_pixmap(alpha=True).save(png)  # alpha PNG is refused by img2pdf -> PyMuPDF path
    out = normalize_to_pdf(str(png), "png", str(tmp_path / "n.pdf"))
    assert out["source_type"] == "scan" and len(out["pages"]) == 1


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice not installed")
def test_docx_through_libreoffice_keeps_text_layer(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("공사도급계약서 계약금액 일금 삼억오천만원정 (₩350,000,000) 계약일자 2026년 3월 5일")
    path = tmp_path / "a.docx"
    doc.save(path)
    out = normalize_to_pdf(str(path), "docx", str(tmp_path / "n.pdf"))
    assert out["source_type"] == "native"
    with fitz.open(out["normalized_pdf_path"]) as pdf:
        assert "350,000,000" in pdf[0].get_text()


def test_garbage_libreoffice_output_is_rejected(tmp_path, monkeypatch):
    def fake_lo(src, out):  # LibreOffice reading an unknown binary as plain text
        doc = fitz.open()
        doc.new_page().insert_text((40, 60), "ĐĎ#ŕˇ±#á####>###ţ˙ ˙˙˙˙˙˙", fontsize=10)
        doc.save(out)
        return True

    monkeypatch.setattr(parser, "libreoffice_to_pdf", fake_lo)
    src = tmp_path / "a.hwpx"
    src.write_bytes(hwpx_bytes())
    out = normalize_to_pdf(str(src), "hwpx", str(tmp_path / "n.pdf"))
    with fitz.open(out["normalized_pdf_path"]) as doc:
        assert "한빛시 도로공사" in doc[0].get_text()
