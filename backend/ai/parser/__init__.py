"""① Format normalization (Document Parser): hwp/hwpx/docx/xlsx/pdf/png/jpg -> normalized PDF.

Every later stage uses only the normalized PDF coordinate system (05_OCR_Pipeline §3-①).
Native formats are also parsed to text/tables (parsed_text.json) with their own parsers;
the PDF itself comes from LibreOffice headless. Without LibreOffice the parsed text is laid
out on A4 pages so the pipeline still runs (layout is lost, text is real).
"""

import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

import pymupdf as fitz

from ai.ocr import classify_page, source_type_of
from ai.parser.hwp import parse_hwp
from ai.types import NormalizedDocument, PageMetaDict

SUPPORTED_FORMATS = {"pdf", "hwp", "hwpx", "docx", "xlsx", "png", "jpg", "jpeg", "doc", "xls"}  # doc/xls: LibreOffice only

# A parsed document is a list of blocks in reading order:
#   {"type": "p", "text": str} | {"type": "table", "cells": [[row, col, text], ...]}  (row/col 1-based)
Block = dict


class UnsupportedFormatError(Exception):
    """Unsupported upload format (API: 415 UNSUPPORTED_FORMAT); the pipeline marks the document FAILED."""


def _local(el) -> str:
    return el.tag.rsplit("}", 1)[-1] if isinstance(el.tag, str) else ""


def parse_hwpx(path: str) -> list[Block]:
    """HWPX (OWPML) = zip of Contents/section*.xml. Paragraph text from <hp:t>, tables from <hp:tbl>/<hp:tc>."""
    from lxml import etree

    blocks: list[Block] = []
    with zipfile.ZipFile(path) as zf:
        names = [n for n in zf.namelist() if re.fullmatch(r"Contents/section\d+\.xml", n)]
        for name in sorted(names, key=lambda n: int(n[len("Contents/section") : -len(".xml")])):  # section10 after section2
            root = etree.fromstring(zf.read(name))
            for p in (el for el in root if _local(el) == "p"):
                text = "".join("".join(t.itertext()) for t in p.iter() if _local(t) == "t" and not _in_table(t, p))
                if text.strip():
                    blocks.append({"type": "p", "text": text})
                for tbl in (el for el in p.iter() if _local(el) == "tbl" and not _in_table(el, p)):
                    cells: list[list] = []
                    for tc in (el for el in tbl.iter() if _local(el) == "tc" and _nearest_table(el) is tbl):
                        addr = next((el for el in tc if _local(el) == "cellAddr"), None)
                        r = int(addr.get("rowAddr", 0)) + 1 if addr is not None else len(cells) + 1
                        c = int(addr.get("colAddr", 0)) + 1 if addr is not None else 1
                        text = " ".join("".join(t.itertext()) for t in tc.iter() if _local(t) == "t").strip()
                        cells.append([r, c, text])
                    blocks.append({"type": "table", "cells": cells})
    return blocks


def _nearest_table(el):
    el = el.getparent()
    while el is not None and _local(el) != "tbl":
        el = el.getparent()
    return el


def _in_table(el, stop) -> bool:
    """True if `el` sits inside a <tbl> below `stop` (the paragraph being read)."""
    el = el.getparent()
    while el is not None and el is not stop:
        if _local(el) == "tbl":
            return True
        el = el.getparent()
    return False


def parse_docx(path: str) -> list[Block]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(path)
    blocks: list[Block] = []
    for child in doc.element.body.iterchildren():
        if _local(child) == "p":
            text = Paragraph(child, doc).text
            if text.strip():
                blocks.append({"type": "p", "text": text})
        elif _local(child) == "tbl":
            seen, cells = set(), []
            for r, row in enumerate(Table(child, doc).rows, start=1):
                for c, cell in enumerate(row.cells, start=1):
                    if id(cell._tc) not in seen:  # merged cells repeat the same <w:tc>
                        seen.add(id(cell._tc))
                        cells.append([r, c, cell.text.strip()])
            blocks.append({"type": "table", "cells": cells})
    return blocks


def parse_xlsx(path: str) -> list[Block]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    blocks: list[Block] = []
    for ws in wb.worksheets:
        blocks.append({"type": "p", "text": f"[{ws.title}]"})
        cells = [
            [r, c, str(v).strip()]
            for r, row in enumerate(ws.iter_rows(values_only=True), start=1)
            for c, v in enumerate(row, start=1)
            if v is not None and str(v).strip()
        ]
        blocks.append({"type": "table", "cells": cells})
    wb.close()
    return blocks


PARSERS = {"hwp": parse_hwp, "hwpx": parse_hwpx, "docx": parse_docx, "xlsx": parse_xlsx}


def blocks_to_lines(blocks: list[Block]) -> list[str]:
    lines: list[str] = []
    for b in blocks:
        if b["type"] == "p":
            lines.extend(t for t in b["text"].splitlines() if t.strip())
            continue
        rows: dict[int, list] = {}
        for r, c, text in b["cells"]:
            rows.setdefault(r, []).append((c, text))
        for r in sorted(rows):
            row = " | ".join(" ".join(t.split()) for _, t in sorted(rows[r]) if t.strip())
            if row:
                lines.append(row)
    return lines


def libreoffice_to_pdf(src: str, out_pdf: str, timeout: int = 120) -> bool:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return False
    with tempfile.TemporaryDirectory() as tmp:
        # private profile dir so concurrent workers don't fight over the user installation lock
        cmd = [exe, f"-env:UserInstallation=file://{tmp}/profile", "--headless", "--convert-to", "pdf", "--outdir", tmp, src]
        try:
            result = subprocess.run(cmd, capture_output=True, timeout=timeout)
        except (subprocess.TimeoutExpired, OSError):
            return False
        produced = Path(tmp) / (Path(src).stem + ".pdf")
        if result.returncode != 0 or not produced.exists():
            return False
        shutil.move(str(produced), out_pdf)
    return True


def conversion_matches(pdf_path: str, blocks: list[Block]) -> bool:
    """LibreOffice opens formats it can't read (e.g. HWP 5.x without H2Orestart) as plain text and still exits 0.
    Accept its PDF only if most text snippets from our own parser appear in it."""
    snippets = [t.replace(" ", "") for b in blocks for t in ([b["text"]] if b["type"] == "p" else [c[2] for c in b["cells"]])]
    snippets = [t[:15] for t in snippets if len(t) >= 4][:30]
    if not snippets:
        return True
    with fitz.open(pdf_path) as doc:
        text = "".join(page.get_text() for page in doc).replace(" ", "").replace("\n", "")
    return sum(t in text for t in snippets) >= 0.5 * len(snippets)


def render_lines_to_pdf(lines: list[str], out_pdf: str) -> None:
    """Fallback without LibreOffice: flow the parsed text onto A4 pages with the built-in CJK font."""
    width, height, margin, size, leading = 595, 842, 50, 10, 16
    font = fitz.Font("korea")
    doc = fitz.open()
    page, y = doc.new_page(width=width, height=height), margin + size
    for line in lines:
        for wrapped in _wrap(line, font, size, width - 2 * margin):
            if y > height - margin:
                page, y = doc.new_page(width=width, height=height), margin + size
            page.insert_text((margin, y), wrapped, fontsize=size, fontname="korea-s")
            y += leading
    doc.save(out_pdf)
    doc.close()


def _wrap(text: str, font, size: float, max_width: float) -> list[str]:
    out, cur = [], ""
    for ch in text:
        if cur and font.text_length(cur + ch, fontsize=size) > max_width:
            out.append(cur)
            cur = ch.lstrip()
        else:
            cur += ch
    return out + [cur] if cur.strip() else out or [""]


def image_to_pdf(src: str, out_pdf: str) -> None:
    import img2pdf

    try:
        Path(out_pdf).write_bytes(img2pdf.convert(src, rotation=img2pdf.Rotation.ifvalid))
    except (img2pdf.AlphaChannelError, ValueError):  # img2pdf refuses alpha PNGs; let PyMuPDF re-encode
        with fitz.open(src) as img:
            Path(out_pdf).write_bytes(img.convert_to_pdf())


def normalize_to_pdf(input_path: str, source_format: str, output_path: str) -> NormalizedDocument:
    fmt = source_format.lower().lstrip(".")
    if fmt not in SUPPORTED_FORMATS:
        raise UnsupportedFormatError(f"'{source_format}' 포맷은 지원하지 않습니다")

    parsed_path = None
    if fmt == "pdf":
        shutil.copyfile(input_path, output_path)
    elif fmt in ("png", "jpg", "jpeg"):
        image_to_pdf(input_path, output_path)
    else:
        blocks: list[Block] | None = None
        parse_error: object = f"no native parser for {fmt}"
        try:
            blocks = PARSERS[fmt](input_path) if fmt in PARSERS else None
        except Exception as exc:  # noqa: BLE001 - a broken native parse falls back to LibreOffice -> scan path
            parse_error = exc
        if blocks is not None:
            parsed_path = str(Path(output_path).with_suffix(".parsed_text.json"))
            Path(parsed_path).write_text(json.dumps(blocks, ensure_ascii=False), encoding="utf-8")
        converted = libreoffice_to_pdf(input_path, output_path)
        if converted and blocks is not None and not conversion_matches(output_path, blocks):
            converted = False
        if not converted:
            if blocks is None:
                raise ValueError(f"{fmt} 변환 실패(LibreOffice 없음/실패): {parse_error}")
            render_lines_to_pdf(blocks_to_lines(blocks), output_path)

    with fitz.open(output_path) as doc:
        if doc.page_count == 0:
            raise ValueError("정규화된 PDF에 페이지가 없습니다")
        pages: list[PageMetaDict] = []
        kinds = []
        for page in doc:
            kinds.append(classify_page(page))
            pages.append(
                {
                    "page_no": page.number + 1,
                    "width_pt": page.rect.width,  # page.rect is already rotation-applied
                    "height_pt": page.rect.height,
                    "rotation": page.rotation,
                    "has_text_layer": bool(page.get_text("text").strip()),
                }
            )
    return {"normalized_pdf_path": output_path, "pages": pages, "source_type": source_type_of(kinds), "parsed_text_path": parsed_path}
